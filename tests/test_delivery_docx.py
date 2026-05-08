import ast
import io
import re
import types
import unittest
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import RGBColor


APP_PATH = Path(__file__).resolve().parents[1] / "app.py"
TARGET_FUNCTIONS = {
    "sanitize_highlighted_article",
    "set_docx_run_font",
    "apply_docx_default_font",
    "append_docx_inline_html",
    "append_highlighted_html_to_docx",
    "append_plain_text_block_to_docx",
    "create_delivery_docx",
}


def load_delivery_helpers():
    source = APP_PATH.read_text(encoding="utf-8")
    tree = ast.parse(source, filename=str(APP_PATH))
    module = types.ModuleType("delivery_docx_helpers")
    module.__dict__.update(
        {
            "io": io,
            "re": re,
            "BeautifulSoup": BeautifulSoup,
            "Document": Document,
            "WD_COLOR_INDEX": WD_COLOR_INDEX,
            "qn": qn,
            "RGBColor": RGBColor,
        }
    )

    for node in tree.body:
        if isinstance(node, ast.FunctionDef) and node.name in TARGET_FUNCTIONS:
            function_module = ast.Module(body=[node], type_ignores=[])
            compiled = compile(function_module, filename=str(APP_PATH), mode="exec")
            exec(compiled, module.__dict__)

    missing = TARGET_FUNCTIONS.difference(module.__dict__)
    if missing:
        raise RuntimeError(f"Missing helper functions: {sorted(missing)}")
    return module


class DeliveryDocxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.helpers = load_delivery_helpers()

    def test_create_delivery_docx_uses_highlighted_html_and_preserves_formatting(self):
        highlighted_html = (
            "<h3>高亮章节</h3>"
            "<p>普通正文 <span class=\"highlight-positive\">重点标注</span> "
            "<span class=\"highlight-risk\">风险提示</span></p>"
        )

        docx_bytes = self.helpers.create_delivery_docx(
            "这段纯正文不应该作为主导出的正文。",
            highlighted_html=highlighted_html,
            script_text=None,
        )
        doc = Document(io.BytesIO(docx_bytes))
        full_text = "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text)

        self.assertIn("高亮阅读版", full_text)
        self.assertIn("高亮章节", full_text)
        self.assertIn("重点标注", full_text)
        self.assertIn("风险提示", full_text)
        self.assertNotIn("这段纯正文不应该作为主导出的正文。", full_text)

        positive_run = None
        risk_run = None
        heading_run = None
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if run.text == "重点标注":
                    positive_run = run
                elif run.text == "风险提示":
                    risk_run = run
                elif run.text == "高亮章节":
                    heading_run = run

        self.assertIsNotNone(positive_run)
        self.assertIsNotNone(risk_run)
        self.assertIsNotNone(heading_run)
        self.assertTrue(positive_run.bold)
        self.assertTrue(risk_run.bold)
        self.assertEqual(str(positive_run.font.color.rgb), "1F57B8")
        self.assertEqual(str(risk_run.font.color.rgb), "B3261E")
        self.assertEqual(positive_run.font.highlight_color, WD_COLOR_INDEX.TURQUOISE)
        self.assertEqual(risk_run.font.highlight_color, WD_COLOR_INDEX.PINK)

        for target_run in (positive_run, risk_run, heading_run):
            east_asia_font = target_run._element.rPr.rFonts.get(qn("w:eastAsia"))
            self.assertEqual(east_asia_font, "微软雅黑")


if __name__ == "__main__":
    unittest.main()
